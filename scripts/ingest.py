"""Download City of Toronto open data and build data/ordis.sqlite.

Usage (from repo root):
    backend/.venv/Scripts/python scripts/ingest.py            # download (if missing) + build
    backend/.venv/Scripts/python scripts/ingest.py --refresh  # re-download everything

Resource URLs are resolved through CKAN package_show by resource name, so a
city-side resource ID change doesn't break ingestion.
"""
import argparse
import csv
import datetime as dt
import json
import sqlite3
import sys
import zipfile
from pathlib import Path

import httpx
import numpy as np
import shapefile  # pyshp
import shapely
from pyproj import Transformer
from shapely.geometry import Polygon, shape

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from backend.data.normalize import normalize_street  # noqa: E402

RAW = ROOT / "data" / "raw"
DB_PATH = ROOT / "data" / "ordis.sqlite"
CKAN = "https://ckan0.cf.opendata.inter.prod-toronto.ca"
TARGET_CRS = 2952  # NAD83(CSRS) / MTM zone 10, metres
# Dev-app X/Y are NAD27 / MTM zone 10. Checked against 3,000 matching Address Points:
# EPSG:7991 median error 9.7 m, EPSG:2019 17.1 m, EPSG:2952 222 m.
DEV_APP_CRS = 7991

csv.field_size_limit(2**31 - 1)

DATASETS = {
    "dev_apps": {
        "package": "development-applications",
        "resource": "Development Applications",  # datastore resource
        "dump_fields": None,
    },
    "parcels": {
        "package": "property-boundaries",
        "resource": "Property Boundaries",
        "dump_fields": "PARCELID,FEATURE_TYPE,STATEDAREA,geometry",
    },
    "address_points": {
        "package": "address-points-municipal-toronto-one-address-repository",
        "resource": "Address Points",
        "dump_fields": "ADDRESS_FULL,LO_NUM,HI_NUM,LINEAR_NAME_FULL,ADDRESS_CLASS_DESC,WARD_NAME,geometry",
    },
    "massing": {
        "package": "3d-massing",
        "resource": "3DMassingShapefile_2025_WGS84.zip",  # file download, not datastore
        "dump_fields": None,
    },
    "trees": {
        "package": "street-tree-data",
        "resource": "Street Tree Data",
        "dump_fields": "COMMON_NAME,DBH_TRUNK,geometry",
    },
    "neighbourhoods": {
        "package": "neighbourhoods",
        "resource": "Neighbourhoods",  # 158-neighbourhood model
        "dump_fields": "AREA_SHORT_CODE,AREA_NAME,geometry",
    },
    "profiles": {
        "package": "neighbourhood-profiles",
        "resource": "neighbourhood-profiles-2021-158-model",  # XLSX download
        "dump_fields": None,
    },
    "heritage": {
        "package": "heritage-register",
        "resource": "Heritage Register Data",  # zipped point shapefile, WGS84
        "dump_fields": None,
    },
    "rentals": {
        "package": "apartment-building-evaluation",
        "resource": "Apartment Building Evaluations 2023 - current.csv",  # RentSafeTO, has lat/lon
        "dump_fields": None,
    },
    "green_spaces": {
        "package": "green-spaces",
        "resource": "Green Spaces - 2952.csv",
        "dump_fields": None,
    },
    "schools": {
        "package": "school-locations-all-types",
        "resource": "School locations-all types data - 2952.csv",
        "dump_fields": None,
    },
    "transit": {
        "package": "ttc-routes-and-schedules",
        "resource": "TTC Routes and Schedules Data",  # GTFS zip
        "dump_fields": None,
    },
}
FILE_EXT = {"massing": "zip", "profiles": "xlsx", "heritage": "zip", "transit": "zip"}
# Tables each dataset writes (used by --only to replace just those).
TABLES = {
    "dev_apps": ["dev_apps", "dev_apps_idx"],
    "parcels": ["parcels", "parcels_idx"],
    "address_points": ["address_points"],
    "massing": ["buildings", "buildings_idx"],
    "trees": ["street_trees", "street_trees_idx"],
    "neighbourhoods": ["neighbourhoods", "neighbourhoods_idx"],
    "profiles": ["nbhd_housing"],
    "heritage": ["heritage", "heritage_idx"],
    "rentals": ["rental_buildings", "rental_buildings_idx"],
    "green_spaces": ["green_spaces", "green_spaces_idx"],
    "schools": ["schools", "schools_idx"],
    "transit": ["transit_stops", "transit_stops_idx"],
}
# Columns added after a database was first built (CREATE TABLE IF NOT EXISTS won't add them).
MIGRATIONS = [("nbhd_housing", "population", "INTEGER"), ("nbhd_housing", "avg_household_size", "REAL")]


# ---------------------------------------------------------------- download

def resolve(key: str) -> dict:
    spec = DATASETS[key]
    r = httpx.get(f"{CKAN}/api/3/action/package_show", params={"id": spec["package"]}, timeout=60)
    r.raise_for_status()
    pkg = r.json()["result"]
    res = next(x for x in pkg["resources"] if x["name"] == spec["resource"])
    if res.get("datastore_active"):
        url = f"{CKAN}/datastore/dump/{res['id']}"
        if spec["dump_fields"]:
            url += f"?fields={spec['dump_fields']}"
    else:
        url = res["url"]
    return {
        "title": pkg["title"],
        "page_url": f"https://open.toronto.ca/dataset/{spec['package']}/",
        "resource_url": url,
        "last_refreshed": pkg.get("last_refreshed"),
    }


def download(url: str, dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_suffix(dest.suffix + ".part")
    print(f"  downloading {url}")
    with httpx.stream("GET", url, timeout=None, follow_redirects=True) as r:
        r.raise_for_status()
        with open(tmp, "wb") as f:
            for chunk in r.iter_bytes(1 << 20):
                f.write(chunk)
    tmp.replace(dest)


def fetch_all(refresh: bool, keys) -> dict:
    meta = {}
    for key in keys:
        info = resolve(key)
        dest = RAW / f"{key}.{FILE_EXT.get(key, 'csv')}"
        if refresh or not dest.exists():
            download(info["resource_url"], dest)
        else:
            print(f"  cached {dest.name}")
        info["downloaded_at"] = dt.datetime.fromtimestamp(dest.stat().st_mtime).isoformat(timespec="seconds")
        meta[key] = (info, dest)
    return meta


# ---------------------------------------------------------------- helpers

def transformer(src: int):
    t = Transformer.from_crs(src, TARGET_CRS, always_xy=True)

    def fn(coords: np.ndarray) -> np.ndarray:
        x, y = t.transform(coords[:, 0], coords[:, 1])
        return np.column_stack([x, y])
    return fn


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def to_int(v):
    f = to_float(v)
    return int(f) if f is not None else None


def proper(v) -> str:
    """QUEEN'S PARK -> Queen's Park (str.title would give Queen'S)."""
    return " ".join(w[:1].upper() + w[1:].lower() for w in str(v or "").split())


def blank_to_none(v):
    v = (v or "").strip()
    return v or None


def insert_geoms(db, table: str, rows: list[tuple], geoms: np.ndarray, cols: str) -> None:
    """rows: attribute tuples aligned with geoms (EPSG:2952). Adds id, geom, and the R*Tree."""
    geoms = shapely.make_valid(geoms)
    bounds = shapely.bounds(geoms)
    wkb = shapely.to_wkb(geoms)
    n_cols = cols.count(",") + 1
    placeholders = ",".join("?" * (n_cols + 2))
    db.executemany(
        f"INSERT INTO {table} (id, {cols}, geom) VALUES ({placeholders})",
        ((i, *row, wkb[i]) for i, row in enumerate(rows)),
    )
    db.executemany(
        f"INSERT INTO {table}_idx VALUES (?,?,?,?,?)",
        ((i, b[0], b[2], b[1], b[3]) for i, b in enumerate(bounds)),
    )


# ---------------------------------------------------------------- loaders

def load_dev_apps(db, path: Path) -> None:
    to_target = Transformer.from_crs(DEV_APP_CRS, TARGET_CRS, always_xy=True)
    to_wgs = Transformer.from_crs(DEV_APP_CRS, 4326, always_xy=True)
    n = 0
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            x0, y0 = to_float(r["X"]), to_float(r["Y"])
            x = y = lon = lat = None
            if x0 and y0:
                x, y = to_target.transform(x0, y0)
                lon, lat = to_wgs.transform(x0, y0)
            street = " ".join(filter(None, [r["STREET_NAME"], r["STREET_TYPE"], r["STREET_DIRECTION"]]))
            rid = int(r["_id"])
            db.execute(
                "INSERT INTO dev_apps VALUES (" + ",".join("?" * 28) + ")",
                (
                    rid, r["APPLICATION#"], r["APPLICATION_TYPE"], blank_to_none(r["STREET_NUM"]),
                    blank_to_none(r["STREET_NAME"]), blank_to_none(r["STREET_TYPE"]),
                    blank_to_none(r["STREET_DIRECTION"]), normalize_street(street), blank_to_none(r["POSTAL"]),
                    blank_to_none(r["DATE_SUBMITTED"]), blank_to_none(r["STATUS"]), blank_to_none(r["DESCRIPTION"]),
                    blank_to_none(r["REFERENCE_FILE#"]), blank_to_none(r["FOLDERRSN"]),
                    blank_to_none(r["WARD_NUMBER"]), blank_to_none(r["WARD_NAME"]),
                    blank_to_none(r["COMMUNITY_MEETING_DATE"]), blank_to_none(r["COMMUNITY_MEETING_TIME"]),
                    blank_to_none(r["COMMUNITY_MEETING_LOCATION"]), blank_to_none(r["APPLICATION_URL"]),
                    blank_to_none(r["CONTACT_NAME"]), blank_to_none(r["CONTACT_PHONE"]),
                    blank_to_none(r["CONTACT_EMAIL"]), blank_to_none(r["PARENT_FOLDER_NUMBER"]), x, y, lon, lat,
                ),
            )
            if x is not None:
                db.execute("INSERT INTO dev_apps_idx VALUES (?,?,?,?,?)", (rid, x, x, y, y))
            n += 1
    print(f"  dev_apps: {n} rows")


def load_parcels(db, path: Path) -> None:
    rows, geoms = [], []
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if not r["geometry"]:
                continue
            geoms.append(shape(json.loads(r["geometry"])))
            rows.append((to_int(r["PARCELID"]), r["FEATURE_TYPE"], r["STATEDAREA"]))
    geoms = shapely.transform(np.array(geoms, dtype=object), transformer(4326))
    insert_geoms(db, "parcels", rows, geoms, "parcel_id, feature_type, stated_area")
    print(f"  parcels: {len(rows)} rows")


def load_address_points(db, path: Path) -> None:
    t = Transformer.from_crs(4326, TARGET_CRS, always_xy=True)
    n = 0
    with open(path, newline="", encoding="utf-8") as f:
        batch = []
        for r in csv.DictReader(f):
            if not r["geometry"]:
                continue
            lon, lat = json.loads(r["geometry"])["coordinates"][:2]
            x, y = t.transform(lon, lat)
            batch.append((
                n, r["ADDRESS_FULL"], to_int(r["LO_NUM"]), to_int(r["HI_NUM"]),
                normalize_street(r["LINEAR_NAME_FULL"]), r["ADDRESS_CLASS_DESC"], r["WARD_NAME"],
                x, y, lon, lat,
            ))
            n += 1
        db.executemany("INSERT INTO address_points VALUES (?,?,?,?,?,?,?,?,?,?,?)", batch)
    print(f"  address_points: {n} rows")


def load_trees(db, path: Path) -> None:
    t = Transformer.from_crs(4326, TARGET_CRS, always_xy=True)
    rows = []
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if not r["geometry"]:
                continue
            g = json.loads(r["geometry"])
            lon, lat = (g["coordinates"][0] if g["type"] == "MultiPoint" else g["coordinates"])[:2]
            x, y = t.transform(lon, lat)
            rows.append((len(rows), r["COMMON_NAME"], to_int(r["DBH_TRUNK"]), x, y))
    db.executemany("INSERT INTO street_trees VALUES (?,?,?,?,?)", rows)
    db.executemany("INSERT INTO street_trees_idx VALUES (?,?,?,?,?)", ((r[0], r[3], r[3], r[4], r[4]) for r in rows))
    print(f"  street_trees: {len(rows)} rows")


def load_neighbourhoods(db, path: Path) -> None:
    rows, geoms = [], []
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r["geometry"]:
                geoms.append(shape(json.loads(r["geometry"])))
                rows.append((int(r["AREA_SHORT_CODE"]), r["AREA_NAME"]))
    geoms = shapely.make_valid(shapely.transform(np.array(geoms, dtype=object), transformer(4326)))
    bounds, wkb = shapely.bounds(geoms), shapely.to_wkb(geoms)
    db.executemany("INSERT INTO neighbourhoods VALUES (?,?,?)", ((c, n, wkb[i]) for i, (c, n) in enumerate(rows)))
    db.executemany("INSERT INTO neighbourhoods_idx VALUES (?,?,?,?,?)",
                   ((rows[i][0], b[0], b[2], b[1], b[3]) for i, b in enumerate(bounds)))
    print(f"  neighbourhoods: {len(rows)} rows")


# Row labels in the 2021 profile workbook (first match wins; leading spaces are stripped).
PROFILE_ROWS = {
    "households": "Total - Private households by tenure - 25% sample data",
    "owners": "Owner",
    "renters": "Renter",
    "apartments_5plus": "Apartment in a building that has five or more storeys",
    "median_value": "Median value of dwellings ($)",
    "median_owner_cost": "Median monthly shelter costs for owned dwellings ($)",
    "median_rent": "Median monthly shelter costs for rented dwellings ($)",
    "pct_owner_30": "% of owner households spending 30% or more of its income on shelter costs",
    "pct_tenant_30": "% of tenant households spending 30% or more of its income on shelter costs",
    "population": "Total - Age groups of the population - 25% sample data",
    "avg_household_size": "Average household size",
}


def load_profiles(db, path: Path) -> None:
    import openpyxl
    rows = list(openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True))
    names, numbers = rows[0][1:], rows[1][1:]
    found = {}
    for r in rows:
        label = str(r[0] or "").strip()
        for key, wanted in PROFILE_ROWS.items():
            if key not in found and label == wanted:
                found[key] = r[1:]
    missing = set(PROFILE_ROWS) - set(found)
    if missing:
        raise RuntimeError(f"profile rows not found: {sorted(missing)}")
    cols = list(PROFILE_ROWS)
    db.executemany(
        f"INSERT INTO nbhd_housing (id, name, {', '.join(cols)}) VALUES ({','.join('?' * (len(cols) + 2))})",
        ((int(numbers[i]), names[i], *[to_float(found[k][i]) for k in cols]) for i in range(len(names))))
    print(f"  nbhd_housing: {len(names)} neighbourhoods")


def load_heritage(db, zip_path: Path) -> None:
    import io
    t = Transformer.from_crs(4326, TARGET_CRS, always_xy=True)
    with zipfile.ZipFile(zip_path) as z:
        stem = next(n for n in z.namelist() if n.lower().endswith(".shp"))[:-4]
        reader = shapefile.Reader(shp=io.BytesIO(z.read(stem + ".shp")), shx=io.BytesIO(z.read(stem + ".shx")),
                                  dbf=io.BytesIO(z.read(stem + ".dbf")))
        rows = []
        for sr in reader.iterShapeRecords():
            if not sr.shape.points:
                continue
            rec = sr.record.as_dict()
            lon, lat = sr.shape.points[0][:2]
            x, y = t.transform(lon, lat)
            text = lambda k: blank_to_none(str(rec.get(k) or ""))  # noqa: E731
            rows.append((len(rows), proper(rec.get("ADDRESS")), text("STATUS"),
                         text("HTG_CONSER"), text("BUILDING_T"), text("DESCRIPTIO"), text("CONSTRUCTI"),
                         text("YEAR_DEMOL"), x, y, lon, lat))
    db.executemany("INSERT INTO heritage VALUES (" + ",".join("?" * 12) + ")", rows)
    db.executemany("INSERT INTO heritage_idx VALUES (?,?,?,?,?)", ((r[0], r[8], r[8], r[9], r[9]) for r in rows))
    print(f"  heritage: {len(rows)} rows")


def load_rentals(db, path: Path) -> None:
    """Latest RentSafeTO evaluation per registered building (RSN)."""
    t = Transformer.from_crs(4326, TARGET_CRS, always_xy=True)
    latest: dict[str, dict] = {}
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            key = (r["YEAR EVALUATED"] or "", r["EVALUATION COMPLETED ON"] or "")
            if r["RSN"] not in latest or key > latest[r["RSN"]]["_key"]:
                latest[r["RSN"]] = r | {"_key": key}
    rows = []
    for r in latest.values():
        lon, lat = to_float(r["LONGITUDE"]), to_float(r["LATITUDE"])
        if lon is None or lat is None:
            continue
        x, y = t.transform(lon, lat)
        rows.append((len(rows), r["RSN"], proper(r["SITE ADDRESS"]), to_int(r["CONFIRMED UNITS"]),
                     to_int(r["CONFIRMED STOREYS"]), to_int(r["YEAR BUILT"]), blank_to_none(r["PROPERTY TYPE"]),
                     to_int(r["CURRENT BUILDING EVAL SCORE"]), blank_to_none(r["EVALUATION COMPLETED ON"]),
                     x, y, lon, lat))
    db.executemany("INSERT INTO rental_buildings VALUES (" + ",".join("?" * 13) + ")", rows)
    db.executemany("INSERT INTO rental_buildings_idx VALUES (?,?,?,?,?)",
                   ((r[0], r[9], r[9], r[10], r[10]) for r in rows))
    print(f"  rental_buildings: {len(rows)} rows")


def load_green_spaces(db, path: Path) -> None:
    rows, geoms = [], []
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if r["geometry"]:
                geoms.append(shape(json.loads(r["geometry"])))  # already EPSG:2952
                rows.append((proper(r["AREA_NAME"]), r["AREA_CLASS"]))
    insert_geoms(db, "green_spaces", rows, np.array(geoms, dtype=object), "name, area_class")
    print(f"  green_spaces: {len(rows)} rows")


def load_schools(db, path: Path) -> None:
    rows = []
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if not r["geometry"]:
                continue
            g = json.loads(r["geometry"])  # EPSG:2952
            x, y = (g["coordinates"][0] if g["type"] == "MultiPoint" else g["coordinates"])[:2]
            board = r["BOARD_NAME"] if r["BOARD_NAME"] not in ("", "None") else None
            rows.append((len(rows), proper(r["NAME"]), blank_to_none(r["SCHOOL_TYPE_DESC"]), board,
                         blank_to_none(r["ADDRESS_FULL"]), x, y))
    db.executemany("INSERT INTO schools VALUES (?,?,?,?,?,?,?)", rows)
    db.executemany("INSERT INTO schools_idx VALUES (?,?,?,?,?)", ((r[0], r[5], r[5], r[6], r[6]) for r in rows))
    print(f"  schools: {len(rows)} rows")


def load_transit(db, zip_path: Path) -> None:
    """GTFS stops plus the modes/routes serving each (via trips and stop_times)."""
    import io
    t = Transformer.from_crs(4326, TARGET_CRS, always_xy=True)
    with zipfile.ZipFile(zip_path) as z:
        def read(name):
            return csv.DictReader(io.TextIOWrapper(z.open(name), "utf-8-sig"))
        routes = {}
        for r in read("routes.txt"):
            kind = {"1": "subway", "3": "bus"}.get(r["route_type"])
            if r["route_type"] == "0":
                kind = "lrt" if r["route_long_name"].startswith("Line") else "streetcar"
            routes[r["route_id"]] = (kind or "other", r["route_short_name"])
        trip_route = {r["trip_id"]: r["route_id"] for r in read("trips.txt")}
        served: dict[str, set] = {}
        reader = csv.reader(io.TextIOWrapper(z.open("stop_times.txt"), "utf-8-sig"))
        head = next(reader)
        ti, si = head.index("trip_id"), head.index("stop_id")
        for row in reader:
            served.setdefault(row[si], set()).add(trip_route.get(row[ti]))
        rows = []
        for s in read("stops.txt"):
            lon, lat = to_float(s["stop_lon"]), to_float(s["stop_lat"])
            serving = [routes[r] for r in served.get(s["stop_id"], ()) if r in routes]
            if lon is None or not serving:
                continue
            x, y = t.transform(lon, lat)
            modes = sorted({m for m, _ in serving})
            pairs = sorted({(n, m) for m, n in serving}, key=lambda p: (len(p[0]), p))
            rows.append((len(rows), s["stop_id"], s["stop_name"], ",".join(modes),
                         ",".join(f"{n}|{m}" for n, m in pairs), x, y))
    db.executemany("INSERT INTO transit_stops VALUES (?,?,?,?,?,?,?)", rows)
    db.executemany("INSERT INTO transit_stops_idx VALUES (?,?,?,?,?)", ((r[0], r[5], r[5], r[6], r[6]) for r in rows))
    print(f"  transit_stops: {len(rows)} rows")


def rings_to_geom(points, parts) -> shapely.Geometry | None:
    """Shapefile rings -> (Multi)Polygon. Outer rings are clockwise, holes counter-clockwise."""
    bounds = list(parts) + [len(points)]
    polys = []
    for a, b in zip(bounds[:-1], bounds[1:]):
        ring = [p[:2] for p in points[a:b]]
        if len(ring) < 4:
            continue
        if shapely.LinearRing(ring).is_ccw and polys:
            polys[-1][1].append(ring)  # hole of the previous shell
        else:
            polys.append((ring, []))
    if not polys:
        return None
    geoms = [Polygon(shell, holes) for shell, holes in polys]
    return geoms[0] if len(geoms) == 1 else shapely.MultiPolygon(geoms)


def load_massing(db, zip_path: Path) -> None:
    out_dir = RAW / "massing"
    with zipfile.ZipFile(zip_path) as z:
        shp_name = next(n for n in z.namelist() if n.lower().endswith(".shp"))
        stem = shp_name[:-4]
        for ext in (".shp", ".shx", ".dbf", ".prj"):
            z.extract(stem + ext, out_dir)
    prj = (out_dir / (stem + ".prj")).read_text()
    # The 2025 "WGS84" file is actually Web Mercator; detect instead of trusting the name.
    src_crs = 3857 if "Mercator" in prj else 4326
    print(f"  massing CRS detected from .prj: EPSG:{src_crs}")

    reader = shapefile.Reader(str(out_dir / stem))
    names = [f[0] for f in reader.fields[1:]]
    rows, geoms, skipped = [], [], 0
    for sr in reader.iterShapeRecords():
        rec = dict(zip(names, sr.record))
        avg = to_float(rec.get("AVG_HEIGHT")) or 0
        msl, surf = to_float(rec.get("HEIGHT_MSL")) or 0, to_float(rec.get("SURF_ELEV")) or 0
        mx = to_float(rec.get("MAX_HEIGHT")) or 0
        # AVG_HEIGHT is populated for every source; MAX_HEIGHT only for LiDAR rows.
        height = avg if avg > 0 else (msl - surf if msl > 0 and surf > 0 and msl > surf else mx)
        geom = rings_to_geom(sr.shape.points, sr.shape.parts)
        if geom is None or height <= 0:
            skipped += 1
            continue
        geoms.append(geom)
        rows.append((round(height, 2), rec.get("HEIGHT_SRC"), rec.get("BLDG_SRC")))
    geoms = shapely.transform(np.array(geoms, dtype=object), transformer(src_crs))
    insert_geoms(db, "buildings", rows, geoms, "height, height_src, bldg_src")
    print(f"  buildings: {len(rows)} rows ({skipped} skipped: no geometry or height)")


# ---------------------------------------------------------------- main

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", action="store_true", help="re-download raw files")
    ap.add_argument("--only", nargs="+", choices=list(DATASETS),
                    help="reload just these datasets into the existing database")
    args = ap.parse_args()
    keys = args.only or list(DATASETS)

    print("Fetching datasets")
    meta = fetch_all(args.refresh, keys)

    if args.only and DB_PATH.exists():
        target = DB_PATH  # update in place
    else:
        target = DB_PATH.with_suffix(".building")
        target.unlink(missing_ok=True)
    db = sqlite3.connect(target)
    db.executescript((ROOT / "backend" / "data" / "schema.sql").read_text())
    for table, column, kind in MIGRATIONS:
        if column not in {r[1] for r in db.execute(f"PRAGMA table_info({table})")}:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {kind}")

    print("Building database")
    loaders = {"dev_apps": load_dev_apps, "address_points": load_address_points,
               "parcels": load_parcels, "massing": load_massing, "trees": load_trees,
               "neighbourhoods": load_neighbourhoods, "profiles": load_profiles, "heritage": load_heritage,
               "rentals": load_rentals, "green_spaces": load_green_spaces, "schools": load_schools,
               "transit": load_transit}
    for key in keys:
        for table in TABLES[key]:
            db.execute(f"DELETE FROM {table}")
        loaders[key](db, meta[key][1])
        info = meta[key][0]
        db.execute(
            "INSERT OR REPLACE INTO sources VALUES (?,?,?,?,?,?)",
            (key, info["title"], info["page_url"], info["resource_url"], info["last_refreshed"], info["downloaded_at"]),
        )
    db.commit()
    db.close()
    if target != DB_PATH:
        target.replace(DB_PATH)
    print(f"Done: {DB_PATH}")


if __name__ == "__main__":
    main()
